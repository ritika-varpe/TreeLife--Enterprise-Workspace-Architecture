import csv
import json
from pathlib import Path

import fitz
import pandas as pd
from docx import Document


SUPPORTED_TYPES = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".csv",
}


def parse_pdf(file_path: str) -> dict:
    pages = []

    with fitz.open(file_path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()

            if text:
                pages.append({
                    "page_number": page_number,
                    "text": text
                })

    return {
        "file_type": "pdf",
        "filename": Path(file_path).name,
        "page_count": len(pages),
        "pages": pages
    }


def parse_docx(file_path: str) -> dict:
    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    tables = []

    for table in document.tables:
        rows = []

        for row in table.rows:
            rows.append([
                cell.text.strip()
                for cell in row.cells
            ])

        tables.append(rows)

    return {
        "file_type": "docx",
        "filename": Path(file_path).name,
        "paragraphs": paragraphs,
        "tables": tables
    }


def parse_excel(file_path: str) -> dict:
    workbook = pd.ExcelFile(file_path)

    sheets = {}

    for sheet_name in workbook.sheet_names:
        dataframe = pd.read_excel(
            file_path,
            sheet_name=sheet_name
        )

        sheets[sheet_name] = {
            "columns": dataframe.columns.tolist(),
            "rows": dataframe.fillna("").to_dict(
                orient="records"
            )
        }

    return {
        "file_type": "xlsx",
        "filename": Path(file_path).name,
        "sheets": sheets
    }


def parse_csv(file_path: str) -> dict:
    dataframe = pd.read_csv(file_path)

    return {
        "file_type": "csv",
        "filename": Path(file_path).name,
        "columns": dataframe.columns.tolist(),
        "rows": dataframe.fillna("").to_dict(
            orient="records"
        )
    }


def parse_file(file_path: str) -> dict:
    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return parse_pdf(file_path)

    if extension == ".docx":
        return parse_docx(file_path)

    if extension == ".xlsx":
        return parse_excel(file_path)

    if extension == ".csv":
        return parse_csv(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


def save_parsed_file(
    file_path: str,
    output_path: str
) -> dict:

    parsed_data = parse_file(file_path)

    output_file = Path(output_path)
    output_file.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    with open(
        output_file,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            parsed_data,
            file,
            indent=2,
            ensure_ascii=False
        )

    return parsed_data